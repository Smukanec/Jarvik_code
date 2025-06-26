import os
import sys
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))
from rag_engine import KnowledgeBase, load_knowledge, search_knowledge


@pytest.fixture
def knowledge_dir(tmp_path):
    """Create a temporary knowledge base with TXT, PDF and DOCX files."""
    folder = tmp_path / "kb"
    folder.mkdir()

    # Always create a TXT file
    (folder / "info.txt").write_text("TXT knowledge", encoding="utf-8")

    # Create a PDF file when the required packages are available
    try:
        from fpdf import FPDF
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        pdf.cell(0, 10, "PDF knowledge")
        pdf.output(str(folder / "info.pdf"))
    except Exception:
        pass

    # Create a DOCX file when the required package is available
    try:
        from docx import Document
        doc = Document()
        doc.add_paragraph("DOCX knowledge")
        doc.save(str(folder / "info.docx"))
    except Exception:
        pass

    return folder


def test_search_knowledge_word_match():
    chunks = [
        "This is a hello text",
        "Another world piece",
        "Completely unrelated",
    ]
    result = search_knowledge("hello world", chunks, threshold=0.9)
    assert result == [
        "Another world piece",
        "This is a hello text",
    ]


def test_search_knowledge_sequence_ratio():
    chunks = ["hello"]
    result = search_knowledge("helo", chunks)
    assert result == ["hello"]


def test_search_knowledge_punctuation_removed():
    chunks = ["hello world", "foo bar"]
    result = search_knowledge("Hello, world!", chunks, threshold=1.1)
    assert result == ["hello world"]


def test_load_knowledge(knowledge_dir):
    chunks = load_knowledge(knowledge_dir)

    assert "TXT knowledge" in chunks

    if any(f.suffix == ".pdf" for f in knowledge_dir.iterdir()):
        assert any("PDF knowledge" in c for c in chunks)

    if any(f.suffix == ".docx" for f in knowledge_dir.iterdir()):
        assert any("DOCX knowledge" in c for c in chunks)


def test_load_knowledge_all_formats(knowledge_dir):
    pytest.importorskip("PyPDF2")
    pytest.importorskip("fpdf")
    pytest.importorskip("docx")

    chunks = load_knowledge(knowledge_dir)

    assert any("PDF knowledge" in c for c in chunks)
    assert any("DOCX knowledge" in c for c in chunks)


def test_knowledge_base_reload(knowledge_dir):
    kb = KnowledgeBase(str(knowledge_dir))
    assert any("TXT knowledge" in c for c in kb.chunks)

    extra = knowledge_dir / "extra.txt"
    extra.write_text("Extra", encoding="utf-8")
    kb.reload()
    assert any("Extra" in c for c in kb.chunks)


def test_search_knowledge_czech_punctuation():
    chunks = [
        "N\u011bco o IPv6 protokolu.",
        "Jin\u00fd text.",
    ]
    result = search_knowledge("n\u011bco o ipv6?", chunks)
    assert result == ["N\u011bco o IPv6 protokolu."]
